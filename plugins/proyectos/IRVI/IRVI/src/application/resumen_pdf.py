# -*- coding: utf-8 -*-
"""
Created on Thu Jun 26 10:21:40 2025

@author: CSANTOS
"""

import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx2pdf import convert

def generar_documento_resumen(df, rut_cliente, nombre_cliente, fecha_operacion,
                              ruta_guardado, imagen_logo):
    """
    Genera un documento Word con un resumen de transacciones.
    """
    nuevo_orden = ['FONDO', 'OPERACION', 'INSTRUMENTO', 'CANTIDAD', 'MONTO']
    df = df.reindex(columns=nuevo_orden)
    # df['CANTIDAD'] = df['CANTIDAD'].apply(lambda x: f"{x:,.0f}".replace(',', '.'))
    # df['MONTO'] = df['MONTO'].apply(lambda x: f"{x:,.0f}".replace(',', '.'))
    df = df.sort_values(by = 'OPERACION')
    document = Document()

    # Encabezado con logo
    header = document.sections[0].header
    header_logo = header.add_paragraph()
    header_logo.alignment = 0
    header_logo.add_run().add_picture(imagen_logo, width=Inches(2))

    # Título y datos del cliente
    document.add_paragraph("Resumen Transacciones", style='Title')
    document.add_paragraph(
        f"Rut Cliente: {rut_cliente}\n\nRazón Social: {nombre_cliente}\n\nFecha Operación: {fecha_operacion}\n\n"
    )

    # Tabla
    table = document.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = col
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(11)
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for _, row_data in df.iterrows():
        row = table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = str(value)
            row.cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            row.cells[j].paragraphs[0].runs[0].font.size = Pt(11)

    ruta_docx = os.path.join(ruta_guardado, f"{nombre_cliente}_resumen.docx")
    document.save(ruta_docx)
    print(f"📄 Documento Word creado: {ruta_docx}")
    ruta_pdf = os.path.join(ruta_guardado, f"ResuTrans{nombre_cliente}.pdf")
    convert(ruta_docx, ruta_pdf)
    print(f"📄 Documento PDF creado: {ruta_pdf}")
    os.remove(ruta_docx)
    