# =============================================================================
                                # OBS
# =============================================================================

# ENCONTRAR TABLA PARA LA CONSULTA DE LOS FOLIOS, CON EL OBJETIVO DE VALIDAR:
# QUE EL ESTADO QUEDE "ENVIADO" Y NO "IMPRESO"

# =============================================================================
                                # LIBRERIAS
# =============================================================================
import tkinter as tk
from tkinter import messagebox
from PIL import Image, ImageTk
import pandas as pd
import os
import sys
import cx_Oracle
sys.path.append(os.path.abspath(r"J:\MOEQUITY\Family Office\Scripts\csantos"))
from accesos import Credentials
from datetime import datetime, timedelta
import xlwings as xw
import time
import shutil
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx2pdf import convert
# =============================================================================
                                # RUTAS
# =============================================================================
RUTA = r"C:\Users\csantos\OneDrive - Banchile\Escritorio\Proyectos\IRVI\DATOS CLIENTES.xlsx"
MACRO_FAC = r"C:\Users\csantos\OneDrive - Banchile\Escritorio\Proyectos\IRVI\Descargar Facturas3.1.xlsm"
DESCARGA = r"C:\Users\csantos\OneDrive - Banchile\Escritorio\Proyectos\IRVI\downloads\archivos"
IMAGEN_FONDO = r"C:\Users\csantos\OneDrive - Banchile\Escritorio\Proyectos\IRVI\downloads\imagenes\Fondo_1.jpeg"
IMAGEN_LOGO = r"C:\Users\csantos\OneDrive - Banchile\Escritorio\Proyectos\IRVI\downloads\imagenes\Logo_1.jpg"
# =============================================================================
                                # VARIABLES
# =============================================================================
# Datos clientes
df = pd.read_excel(RUTA)

# Usuarios
usuarios = {
            "CSANTOS": "Christian Santos",
            "GCASTILL": "Gregori Castillo",
            "JSEPULVE": "Jean Sepulveda",
            "BBRUNA": "Bastian Bruna",
            "JCHANDIA": "Jessica Chandia"
            }
# Lista para etiquetas dinámicas
archivo_labels = []
# Variable global para el botón de ejecución
btn_ejecutar = None
# Fecha manual

# =============================================================================
                                # FUNCIONES
# =============================================================================
def limpiar_carpeta_cliente(ruta_carpeta):
    """
    Elimina todos los archivos dentro de la carpeta del cliente.
    """
    if not os.path.exists(ruta_carpeta):
        print(f"⚠️ La carpeta no existe: {ruta_carpeta}")
        return

    for archivo in os.listdir(ruta_carpeta):
        ruta_archivo = os.path.join(ruta_carpeta, archivo)
        try:
            if os.path.isfile(ruta_archivo) or os.path.islink(ruta_archivo):
                os.unlink(ruta_archivo)
            elif os.path.isdir(ruta_archivo):
                shutil.rmtree(ruta_archivo)
        except Exception as e:
            print(f"❌ Error al eliminar {ruta_archivo}: {e}")

    print(f"🧹 Carpeta limpiada: {ruta_carpeta}")
    
def crear_carpeta_cliente(nombre_cliente, carpeta_base):
    """
    Crea una carpeta con el nombre del cliente dentro de la carpeta base.
    Si ya existe, no hace nada.
    Retorna la ruta completa de la carpeta creada o existente.
    """
    ruta_carpeta = os.path.join(carpeta_base, nombre_cliente)

    try:
        os.makedirs(ruta_carpeta, exist_ok=True)
        print(f"📁 Carpeta lista: {ruta_carpeta}")
    except Exception as e:
        print(f"❌ Error al crear carpeta: {e}")
        ruta_carpeta = None

    return ruta_carpeta

def conectar_bbdd():
    '''Conectar Credenciales'''
    credenciales = Credentials()
    result = credenciales.get_generic_credential('OperacionesBD')
    usuario = result.username
    passw = result.password
    dsn = cx_Oracle.makedsn("standby-bd.bantcent.cl", 1531, service_name="ORA9")
    connection = cx_Oracle.connect(user=usuario, password=passw ,dsn=dsn,encoding="UTF-8")
    cursor = connection.cursor()
    return cursor, connection, usuario, passw

def buscar_cliente():
    ''' Busca cliente e indica los archivos que lleva. '''
    
    global archivo_labels, btn_ejecutar
    id_cliente = entry_id.get().strip()
    resultado = df[df['RUT'].astype(str) == str(id_cliente)]

    # Eliminar campos anteriores para no pisar info.
    for widget in archivo_labels:
        widget.destroy()
    archivo_labels = []
    
    # Eliminar botón anterior si existe.
    if btn_ejecutar:
        btn_ejecutar.destroy()
        btn_ejecutar = None
    
    # Agregar a cliente si corresponde.
    if not resultado.empty:
        nombre = resultado.iloc[0]['CLIENTE']
        entry_resultado.config(state="normal")
        entry_resultado.delete(0, tk.END)
        entry_resultado.insert(0, nombre)
        entry_resultado.config(state="readonly")
        
        # Archivos disponible según cliente.
        archivos_disponibles = []
        campos = ['FACTURA', 'XML', 'RESUMEN', 'TXT', 'Excel Macro']
        y_pos = 230
        for campo in campos:
            valor = str(resultado.iloc[0][campo]).strip().lower()
            if valor == 'si':
                entry_archivo = tk.Entry(ventana, font=("Helvetica", 12), width=60, readonlybackground="white", fg="black")
                entry_archivo.insert(0, f"{campo.capitalize()}: Sí")
                entry_archivo.config(state="readonly")
                canvas.create_window(100, y_pos, window=entry_archivo, anchor="w")
                archivo_labels.append(entry_archivo)
                archivos_disponibles.append(campo)
                y_pos += 35

        # Botón de ejecución de archivos.
        if archivos_disponibles:
            btn_ejecutar = tk.Button(ventana, text="Obtener Archivos Cliente", command=lambda: obtener_archivos(id_cliente, archivos_disponibles))
            canvas.create_window(325, 150, window=btn_ejecutar, anchor="w")
            
        print(f"RUT: {id_cliente}")
        print("Archivos disponibles:")
        for archivo in archivos_disponibles:
            print(f"- {archivo}")
        
        # Retorna cliente y archivos para poder descargar.
        return id_cliente, archivos_disponibles
    else:
        entry_resultado.config(state="normal")
        entry_resultado.delete(0, tk.END)
        entry_resultado.insert(0, "ID no encontrado")
        entry_resultado.config(state="readonly")
        messagebox.showwarning("Aviso", "El ID ingresado no existe.")
        return None, []
    
def qry_folios_fecha(connection, rut_completo, fecha_str):
    ''' Consulta SQL para cliente en ejecución. '''
    # Quitar último dígito
    rut_truncado = rut_completo[:-1]
    sql = """
    SELECT  TG_SRU_TG_CLI_TG_PER_RUT_PER AS Rut,
            TG_SRU_SRU_SRU AS Sub_Rut,
            TO_CHAR(FEC_FAC, 'DD-MM-YYYY') AS Fecha_Operacion,
            FOL_FAC AS Folio
    FROM    ta_fac
    WHERE   FEC_FAC = TO_DATE(:1, 'DD-MM-YYYY')
            AND FOL_FAC IS NOT NULL
            AND TG_SRU_TG_CLI_TG_PER_RUT_PER = :2
            AND USR_IMP_FAC IN ('JCHANDIA','JSEPULVE',
                                'GCASTILL','MCARVACH1',
                                'BBRUNA','CSANTOS')
    """
    try:
        df = pd.read_sql(sql, connection, params=(fecha_str, rut_truncado))
        print(f"\n📋 Resultados para RUT {rut_truncado} y fecha {fecha_str}:")
        if df.empty:
            print("No se encontraron registros.")
        else:
            print(df.to_string(index=False))
        return df
    except Exception as e:
        print(f"❌ Error al ejecutar la consulta: {e}")
        return pd.DataFrame()

def procesar_excel_con_macro(df_resultado, nombre_macro, ruta_descarga):
    if df_resultado.empty:
        print("⚠️ No hay datos para procesar en Excel.")
        return

    print(f"\n📥 Abriendo Excel y ejecutando macro: {nombre_macro}...")

    app = xw.App(visible=False, add_book=False)
    wb = app.books.open(MACRO_FAC)

    macro_limpiar = wb.app.macro("Hoja1.limpiar")
    macro_limpiar()
    time.sleep(1)

    sheet = wb.sheets[0]
    for cell in range(len(df_resultado)):
        sheet.range(cell + 6, 1).value = df_resultado['RUT'].iloc[cell]
        sheet.range(cell + 6, 2).value = df_resultado['SUB_RUT'].iloc[cell]
        sheet.range(cell + 6, 3).value = df_resultado['FOLIO'].iloc[cell]
        sheet.range(cell + 6, 4).value = 33
        sheet.range(cell + 6, 5).value = ruta_descarga

    print(f"\n📄 Ejecutando macro para {nombre_macro}...")
    time.sleep(1)
    macro = wb.app.macro(f"Hoja1.{nombre_macro}")
    macro()
    time.sleep(1)

    wb.close()
    app.quit()
    print("✅ Proceso completado.")

def qry_resumen_cliente(connection, rut_completo, fecha_str):
    """
    Consulta los datos de resumen para un cliente específico y una fecha dada.
    """
    rut_truncado = rut_completo[:-1]  # Eliminar el último dígito

    sql = """
SELECT 
        sru.cnp_sru Fondo, 
        decode(OBN.tip_ope_ord,'VTA','V','C') Operacion,  
        OBN.ta_ser_nom_ser Instrumento,  
        cg_fnc_gbl.fg_fmt_num(sum(ASG.cnt_asg),0,0) Cantidad, 
        cg_fnc_gbl.fg_fmt_num(trunc(sum(ASG.cnt_asg * ASG.pre_asg)),0,0) Monto, 
        to_char(asg.fec_asg,'dd-mm-yyyy') Fecha_trade, 
        to_char(trs.fec_liq_trs,'dd-mm-yyyy')  Fecha_Pago 
            
FROM
        ta_obn obn, 
        ta_asg asg, 
        ta_ult_trs trs, 
        tg_sru sru
         
WHERE
        obn.num_ord = asg.ta_obn_num_ord 
        AND asg.fec_asg >= TO_DATE(:1, 'DD-MM-YYYY')
        and obn.tg_SRU_TG_CLI_TG_per_rut_per = :2
        and trs.cor_trs = asg.ta_ult_trs_cor_trs 
        and SRU.TG_CLI_TG_PER_RUT_PER = OBN.TG_SRU_TG_CLI_TG_PER_RUT_PER 
        and SRU.SRU_SRU = OBN.TG_SRU_SRU_SRU 
        
    group by sru.cnp_sru, 
             to_char(asg.fec_asg,'dd-mm-yyyy'), 
             decode(OBN.tip_ope_ord,'VTA','V','C') , 
             OBN.ta_ser_nom_ser, 
             obn.tg_sru_sru_sru, 
             to_char(trs.fec_liq_trs,'dd-mm-yyyy')
    """

    try:
        df = pd.read_sql(sql, connection, params=(fecha_str, rut_truncado))
        print(f"\n📋 Resultados de resumen para RUT {rut_truncado} y fecha {fecha_str}:")
        if df.empty:
            print("No se encontraron registros de resumen.")
        else:
            print(df.to_string(index=False))
        return df
    except Exception as e:
        print(f"❌ Error al ejecutar la consulta de resumen: {e}")
        return pd.DataFrame()

def generar_documento_resumen(df, rut_cliente, nombre_cliente, fecha_operacion, ruta_guardado):
    """
    Genera un documento Word con un resumen de transacciones.
    """
    nuevo_orden = ['FONDO', 'OPERACION', 'INSTRUMENTO', 'CANTIDAD', 'MONTO']
    df = df.reindex(columns=nuevo_orden)
    # df['CANTIDAD'] = df['CANTIDAD'].apply(lambda x: f"{x:,.0f}".replace(',', '.'))
    # df['MONTO'] = df['MONTO'].apply(lambda x: f"{x:,.0f}".replace(',', '.'))

    document = Document()

    # Encabezado con logo
    header = document.sections[0].header
    header_logo = header.add_paragraph()
    header_logo.alignment = 0
    header_logo.add_run().add_picture(IMAGEN_LOGO, width=Inches(2))

    # Título y datos del cliente
    document.add_paragraph("Resumen Transacciones", style='Title')
    document.add_paragraph(
        f"Rut Cliente: {rut_cliente}\n\nRazón Social: {nombre_cliente}\n\nFecha Operación: {fecha_operacion}\n\n"
    )

    # Tabla
    table = document.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = col
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(11)
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for _, row_data in df.iterrows():
        row = table.add_row()
        for j, value in enumerate(row_data):
            row.cells[j].text = str(value)
            row.cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            row.cells[j].paragraphs[0].runs[0].font.size = Pt(11)

    ruta_docx = os.path.join(ruta_guardado, f"{nombre_cliente}_resumen.docx")
    document.save(ruta_docx)
    
    print(f"📄 Documento Word creado: {ruta_docx}")
    
    ruta_pdf = os.path.join(ruta_guardado, f"{nombre_cliente}_resumen.pdf")
    convert(ruta_docx, ruta_pdf)
    print(f"📄 Documento PDF creado: {ruta_pdf}")

def obtener_archivos(id_cliente, archivos_disponibles):
    print(f"\n🔍 Obteniendo archivos para cliente {id_cliente}...")

    try:
        cursor, connection, _, _ = conectar_bbdd()
    except Exception as e:
        print(f"❌ Error de conexión: {e}")
        return

    # Asegurar que el RUT y el DataFrame estén en formato string
    id_cliente = str(id_cliente).strip()
    df['RUT'] = df['RUT'].astype(str)

    try:
        nombre_cliente = df[df['RUT'] == id_cliente]['CLIENTE'].values[0]
    except IndexError:
        print("❌ No se encontró el nombre del cliente en el DataFrame.")
        return

    # Crear y limpiar carpeta del cliente solo una vez
    ruta_descarga = crear_carpeta_cliente(nombre_cliente,DESCARGA)
    limpiar_carpeta_cliente(ruta_descarga)

    # Consultar datos desde la base
    fecha_hoy = datetime.today().strftime('%d-%m-%Y')
    fecha_ayer = (datetime.today() - timedelta(days=1)).strftime('%d-%m-%Y')

    df_resultado = qry_folios_fecha(connection, id_cliente, fecha_hoy)
    if df_resultado.empty:
        df_resultado = qry_folios_fecha(connection, id_cliente, fecha_ayer)

    if df_resultado.empty:
        print("⚠️ No se encontraron datos para procesar.")
        return

    # Procesar según tipo de archivo
    for archivo in archivos_disponibles:
        print(f"📦 Procesando {archivo} para {id_cliente}")
        if archivo == "FACTURA":
            procesar_excel_con_macro(df_resultado, "descargarPDF", ruta_descarga)
        elif archivo == "XML":
            procesar_excel_con_macro(df_resultado, "descargarXML", ruta_descarga)

    # 🔄 Procesar RESUMEN si está disponible
    if "RESUMEN" in archivos_disponibles:
        df_resumen = qry_resumen_cliente(connection, id_cliente, fecha_hoy)
        if df_resumen.empty:
            df_resumen = qry_resumen_cliente(connection, id_cliente, fecha_ayer)

        if df_resumen.empty:
            print("⚠️ No se encontraron datos de resumen.")
        else:
            print("✅ Datos de resumen obtenidos correctamente.")
            generar_documento_resumen(
                df_resumen,
                rut_cliente=id_cliente,
                nombre_cliente=nombre_cliente,
                fecha_operacion=fecha_hoy,
                # imagen_logo=IMAGEN_LOGO,
                ruta_guardado=ruta_descarga
            )

# =============================================================================
                                # INTERFAZ
# =============================================================================
# Crear ventana principal
ventana = tk.Tk()
ventana.title("Interfaz Renta Variable Institucionales")
ventana.geometry("800x450")

# Imagen de fondo
canvas = tk.Canvas(ventana, width=800, height=450)
canvas.pack(fill="both", expand=True)

# Cargar imagen de fondo
imagen = Image.open(IMAGEN_FONDO)

imagen = imagen.resize((800, 450), Image.ANTIALIAS)
imagen_fondo = ImageTk.PhotoImage(imagen)
canvas.create_image(0, 0, image=imagen_fondo, anchor="nw")

# Login
usuario_id = os.getlogin()
usuario_nombre = usuarios.get(usuario_id, "Usuario")
canvas.create_text(400, 50, text=f"Bienvenido, {usuario_nombre}.", font=("Helvetica", 16), fill="white", anchor="center")

# Entrada de ID
canvas.create_text(100, 150, text="Rut:", font=("Helvetica", 14), fill="white", anchor="e")
entry_id = tk.Entry(ventana)
canvas.create_window(100, 150, window=entry_id, anchor="w")

# Botón de búsqueda
btn_buscar = tk.Button(ventana, text="Buscar", command=buscar_cliente)
canvas.create_window(250, 150, window=btn_buscar, anchor="w")

# Etiqueta "Cliente:"
canvas.create_text(100, 190, text="Cliente:", font=("Helvetica", 14), fill="white", anchor="e")

entry_resultado = tk.Entry(ventana, font=("Helvetica", 12), width=60, state="readonly", readonlybackground="white", fg="black")
canvas.create_window(100, 190, window=entry_resultado, anchor="w")

# Etiqueta "Archivos:"
canvas.create_text(100, 230, text="Archivos:", font=("Helvetica", 14), fill="white", anchor="e")
# Botón ejecutar
# btn_buscar = tk.Button(ventana, text="Obtener Archivos Cliente", command=buscar_cliente)
# canvas.create_window(100, 370, window=btn_buscar, anchor="w")
ventana.mainloop()