# -*- coding: utf-8 -*-
"""
Created on Thu Jan 23 10:03:44 2025

@author: BBRUNA
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class Asignacion:
    indice: int
    rut: str
    fondo: str
    tipo: str
    nemo: str
    cantidad_objetivo: int
    monto_final_objetivo: int


@dataclass(frozen=True)
class LineaFactura:
    indice: int
    folio: int
    precio_unitario: Decimal
    precio_texto: str
    cantidad: int
    monto_neto: int


@dataclass(frozen=True)
class ResultadoProceso:
    ruta_planilla: Path
    carpeta_archivos_carga: Path
    cantidad_folios: int
    cantidad_lineas_factura: int
    cantidad_asignaciones: int


_RE_ESPACIOS = re.compile(r"\s+")
_RE_NO_NUM = re.compile(r"[^\d,.\-]")

COLUMNAS_ASIGNACIONES = ["rut", "fondo", "tipo", "nemo", "cantidad", "monto_final"]
COLUMNAS_FACTURAS = ["folio", "precio_unitario", "cantidad", "monto_neto"]


def normalizar_texto(valor: object) -> str:
    if pd.isna(valor):
        return ""
    return _RE_ESPACIOS.sub(" ", str(valor)).strip()


def parsear_entero_chileno(valor: object) -> int:
    s = normalizar_texto(valor)
    if not s:
        raise ValueError("Entero vacio/no informado.")
    s = _RE_NO_NUM.sub("", s)
    if "," in s:
        s = s.split(",", 1)[0]
    s = s.replace(".", "")
    if s in {"", "-", "--"}:
        raise ValueError(f"Entero invalido: {valor!r}")
    return int(s)


def parsear_decimal_chileno(valor: object) -> Decimal:
    s = normalizar_texto(valor)
    if not s:
        raise ValueError("Decimal vacio/no informado.")

    s_limpio = _RE_NO_NUM.sub("", s)
    if "," in s_limpio:
        s_limpio = s_limpio.replace(".", "")
        s_limpio = s_limpio.replace(",", ".")
    else:
        if s_limpio.count(".") > 1:
            s_limpio = s_limpio.replace(".", "")

    try:
        return Decimal(s_limpio)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Decimal invalido: {valor!r}") from exc


def formatear_precio_chileno(precio: Decimal) -> str:
    q = precio.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{q:.2f}".replace(".", ",")


def normalizar_tipo(valor: object) -> str:
    s = normalizar_texto(valor).lower()
    if s in {"compra", "c"}:
        return "Compra"
    if s in {"venta", "v"}:
        return "Venta"
    raise ValueError(f"Tipo invalido (esperado Compra/Venta): {valor!r}")


def normalizar_rut_sin_puntos_con_guion(rut: str) -> str:
    return normalizar_texto(rut).replace(".", "").replace(" ", "")


def fondo_como_numero(fondo: str) -> str:
    s = normalizar_texto(fondo)
    if not s:
        return ""
    try:
        n = int(float(s))
        return str(n)
    except Exception:
        return s


def _validar_columnas(df: pd.DataFrame, requeridas: Sequence[str], nombre_hoja: str) -> None:
    faltantes = [c for c in requeridas if c not in df.columns]
    if faltantes:
        raise ValueError(f"En hoja '{nombre_hoja}' faltan columnas: {faltantes}")


def _parsear_asignaciones_df(df_asig: pd.DataFrame) -> List[Asignacion]:
    _validar_columnas(df_asig, COLUMNAS_ASIGNACIONES, "asignaciones")
    asignaciones: List[Asignacion] = []
    for i, row in df_asig.iterrows():
        asignaciones.append(
            Asignacion(
                indice=int(i),
                rut=normalizar_texto(row["rut"]),
                fondo=normalizar_texto(row["fondo"]),
                tipo=normalizar_tipo(row["tipo"]),
                nemo=normalizar_texto(row["nemo"]),
                cantidad_objetivo=parsear_entero_chileno(row["cantidad"]),
                monto_final_objetivo=parsear_entero_chileno(row["monto_final"]),
            )
        )
    return asignaciones


def _parsear_facturas_df(df_fac: pd.DataFrame) -> List[LineaFactura]:
    _validar_columnas(df_fac, COLUMNAS_FACTURAS, "facturas")
    lineas: List[LineaFactura] = []
    for i, row in df_fac.iterrows():
        folio = parsear_entero_chileno(row["folio"])
        precio_dec = parsear_decimal_chileno(row["precio_unitario"])
        precio_texto = formatear_precio_chileno(precio_dec)
        lineas.append(
            LineaFactura(
                indice=int(i),
                folio=folio,
                precio_unitario=precio_dec,
                precio_texto=precio_texto,
                cantidad=parsear_entero_chileno(row["cantidad"]),
                monto_neto=parsear_entero_chileno(row["monto_neto"]),
            )
        )
    return lineas


def leer_excel_entrada(ruta_excel: Path) -> Tuple[List[Asignacion], List[LineaFactura]]:
    if not ruta_excel.exists():
        raise FileNotFoundError(f"No existe el archivo: {ruta_excel}")

    xls = pd.ExcelFile(ruta_excel)
    if "asignaciones" not in xls.sheet_names:
        raise ValueError("Falta la hoja 'asignaciones'.")
    if "facturas" not in xls.sheet_names:
        raise ValueError("Falta la hoja 'facturas'.")

    df_asig = pd.read_excel(ruta_excel, sheet_name="asignaciones", dtype=object)
    df_fac = pd.read_excel(ruta_excel, sheet_name="facturas", dtype=object)

    asignaciones = _parsear_asignaciones_df(df_asig)
    lineas = _parsear_facturas_df(df_fac)
    validar_consistencia(asignaciones, lineas)
    return asignaciones, lineas


def validar_consistencia(asignaciones: Sequence[Asignacion], lineas: Sequence[LineaFactura]) -> None:
    if not asignaciones:
        raise ValueError("No hay filas en 'asignaciones'.")
    if not lineas:
        raise ValueError("No hay filas en 'facturas'.")

    total_asig = sum(a.cantidad_objetivo for a in asignaciones)
    total_fac = sum(l.cantidad for l in lineas)
    if total_asig != total_fac:
        raise ValueError(
            "La suma de cantidades no cuadra:\n"
            f"  asignaciones={total_asig}\n"
            f"  facturas={total_fac}"
        )

    if any(l.folio <= 0 for l in lineas):
        raise ValueError("Hay folios invalidos (<=0).")
    if any(l.cantidad < 0 for l in lineas):
        raise ValueError("Hay lineas con cantidad negativa (no valido).")
    if any(l.monto_neto < 0 for l in lineas):
        raise ValueError("Hay 'monto_neto' negativo en 'facturas' (no esperado).")


def construir_asignacion_cantidades(
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
) -> List[List[int]]:
    n = len(asignaciones)
    m = len(lineas)
    x = [[0] * m for _ in range(n)]
    saldo_fila = [a.cantidad_objetivo for a in asignaciones]
    i = 0

    for j, linea in enumerate(lineas):
        restante_col = linea.cantidad
        while restante_col > 0:
            while i < n and saldo_fila[i] == 0:
                i += 1
            if i >= n:
                raise RuntimeError("Error interno: no hay saldo de fila para completar columna.")
            t = min(saldo_fila[i], restante_col)
            x[i][j] += t
            saldo_fila[i] -= t
            restante_col -= t

    if any(s != 0 for s in saldo_fila):
        raise RuntimeError("Error interno: quedaron saldos de fila no asignados.")
    return x


def repartir_neto_por_linea(
    x: Sequence[Sequence[int]],
    lineas: Sequence[LineaFactura],
) -> List[List[int]]:
    n = len(x)
    neto = [[0] * len(lineas) for _ in range(n)]

    for j, linea in enumerate(lineas):
        total_q = sum(x[i][j] for i in range(n))
        if total_q != linea.cantidad:
            raise RuntimeError("Error interno: suma de columna != cantidad de la linea.")
        if total_q == 0:
            continue

        monto = linea.monto_neto
        asignado_base = 0
        cuotas: List[Tuple[int, int, float]] = []
        for i in range(n):
            q = x[i][j]
            if q == 0:
                continue
            ideal = (monto * q) / total_q
            base = int(math.floor(ideal))
            frac = float(ideal - base)
            cuotas.append((i, base, frac))
            asignado_base += base

        residuo = monto - asignado_base
        cuotas.sort(key=lambda t: t[2], reverse=True)
        for k, (i, base, _) in enumerate(cuotas):
            neto[i][j] = base + (1 if k < residuo else 0)

        if sum(neto[i][j] for i in range(n)) != monto:
            raise RuntimeError("Error interno: no se pudo cuadrar monto_neto por linea.")

    return neto


def calcular_neto_y_comision_total(
    asignaciones: Sequence[Asignacion],
    neto: Sequence[Sequence[int]],
) -> Tuple[List[int], List[int]]:
    neto_total = [sum(neto[i]) for i in range(len(asignaciones))]
    comision_total = [
        asignaciones[i].monto_final_objetivo - neto_total[i] for i in range(len(asignaciones))
    ]
    return neto_total, comision_total


def violaciones_signo(
    asignaciones: Sequence[Asignacion],
    comision_total: Sequence[int],
) -> List[int]:
    malos: List[int] = []
    for i, a in enumerate(asignaciones):
        c = int(comision_total[i])
        if a.tipo == "Compra" and c < 0:
            malos.append(i)
        elif a.tipo == "Venta" and c > 0:
            malos.append(i)
    return malos


def calcular_rate_scaled(lineas: Sequence[LineaFactura], escala: int = 1000) -> List[int]:
    rates: List[int] = []
    for l in lineas:
        if l.cantidad <= 0:
            rates.append(0)
        else:
            rates.append(int(round((l.monto_neto * escala) / l.cantidad)))
    return rates


def aproximar_neto_total_scaled(x: Sequence[Sequence[int]], rate_scaled: Sequence[int]) -> List[int]:
    neto_scaled = [0] * len(x)
    for i in range(len(x)):
        s = 0
        for j in range(len(rate_scaled)):
            if x[i][j]:
                s += rate_scaled[j] * x[i][j]
        neto_scaled[i] = s
    return neto_scaled


def escoger_contraparte(
    asignaciones: Sequence[Asignacion],
    comision_total: Sequence[int],
    idx_candidatos: Sequence[int],
    direccion: str,
) -> int:
    mejor = None
    mejor_score = -10**18
    for k in idx_candidatos:
        a = asignaciones[k]
        c = int(comision_total[k])
        score = 0
        if direccion == "subir_neto_k":
            if a.tipo == "Venta" and c > 0:
                score += 10_000_000
            if a.tipo == "Venta":
                score += 1_000_000
            if a.tipo == "Compra" and c > 0:
                score += min(c, 5_000_000)
        else:
            if a.tipo == "Compra" and c < 0:
                score += 10_000_000
            if a.tipo == "Compra":
                score += 1_000_000
            if a.tipo == "Venta" and c < 0:
                score += min(abs(c), 5_000_000)
        score += (1000 - k)
        if score > mejor_score:
            mejor_score = score
            mejor = k
    return int(mejor) if mejor is not None else -1


def aplicar_swaps_para_signo(
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
    x: List[List[int]],
    max_iteraciones: int = 50_000,
) -> None:
    n = len(asignaciones)
    m = len(lineas)
    rate_scaled = calcular_rate_scaled(lineas, escala=1000)
    cols_alta = sorted(range(m), key=lambda j: rate_scaled[j], reverse=True)
    cols_baja = sorted(range(m), key=lambda j: rate_scaled[j])

    iteracion = 0
    while iteracion < max_iteraciones:
        neto = repartir_neto_por_linea(x, lineas)
        _, comision_total = calcular_neto_y_comision_total(asignaciones, neto)
        malos = violaciones_signo(asignaciones, comision_total)
        if not malos:
            return

        malos.sort(key=lambda i: abs(int(comision_total[i])), reverse=True)
        hizo_movimiento = False
        _ = aproximar_neto_total_scaled(x, rate_scaled)

        for i in malos:
            a_i = asignaciones[i]
            c_i = int(comision_total[i])

            if a_i.tipo == "Compra" and c_i < 0:
                j_alta = next((j for j in cols_alta if x[i][j] > 0), None)
                if j_alta is None:
                    continue
                for j_baja in cols_baja:
                    if rate_scaled[j_baja] >= rate_scaled[j_alta]:
                        break
                    candidatos = [k for k in range(n) if k != i and x[k][j_baja] > 0]
                    if not candidatos:
                        continue
                    k = escoger_contraparte(asignaciones, comision_total, candidatos, "subir_neto_k")
                    if k < 0:
                        continue
                    x[i][j_alta] -= 1
                    x[i][j_baja] += 1
                    x[k][j_baja] -= 1
                    x[k][j_alta] += 1
                    hizo_movimiento = True
                    break

            elif a_i.tipo == "Venta" and c_i > 0:
                j_baja = next((j for j in cols_baja if x[i][j] > 0), None)
                if j_baja is None:
                    continue
                for j_alta in cols_alta:
                    if rate_scaled[j_alta] <= rate_scaled[j_baja]:
                        break
                    candidatos = [k for k in range(n) if k != i and x[k][j_alta] > 0]
                    if not candidatos:
                        continue
                    k = escoger_contraparte(asignaciones, comision_total, candidatos, "bajar_neto_k")
                    if k < 0:
                        continue
                    x[i][j_baja] -= 1
                    x[i][j_alta] += 1
                    x[k][j_alta] -= 1
                    x[k][j_baja] += 1
                    hizo_movimiento = True
                    break

            if hizo_movimiento:
                break

        iteracion += 1
        if not hizo_movimiento:
            neto = repartir_neto_por_linea(x, lineas)
            _, comision_total = calcular_neto_y_comision_total(asignaciones, neto)
            malos = violaciones_signo(asignaciones, comision_total)
            detalles = []
            for i in malos[:15]:
                rut = normalizar_rut_sin_puntos_con_guion(asignaciones[i].rut)
                fondo = fondo_como_numero(asignaciones[i].fondo)
                tipo = asignaciones[i].tipo
                detalles.append(f"{rut} | fondo {fondo} | {tipo} | comision={int(comision_total[i])}")
            msg = "No se pudo corregir el signo de comision para todos.\n" + "\n".join(detalles)
            raise ValueError(msg)

    raise ValueError("Limite de iteraciones alcanzado intentando corregir signos de comision.")


def indices_lineas_por_folio(lineas: Sequence[LineaFactura]) -> Dict[int, List[int]]:
    por_folio: Dict[int, List[int]] = {}
    for j, l in enumerate(lineas):
        por_folio.setdefault(l.folio, []).append(j)
    return por_folio


def orden_precios_input_en_folio(lineas: Sequence[LineaFactura], indices_folio: Sequence[int]) -> List[str]:
    vistos: set[str] = set()
    orden: List[str] = []
    for j in indices_folio:
        p = lineas[j].precio_texto
        if p not in vistos:
            vistos.add(p)
            orden.append(p)
    return orden


def mapa_indices_por_precio_en_folio(
    lineas: Sequence[LineaFactura],
    indices_folio: Sequence[int],
) -> Dict[str, List[int]]:
    mapa: Dict[str, List[int]] = {}
    for j in indices_folio:
        mapa.setdefault(lineas[j].precio_texto, []).append(j)
    return mapa


def elegir_folio_comision_por_cliente(
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
    neto: Sequence[Sequence[int]],
    comision_total: Sequence[int],
) -> Dict[int, int]:
    por_folio = indices_lineas_por_folio(lineas)
    folio_por_cliente: Dict[int, int] = {}
    for i in range(len(asignaciones)):
        netos_folio = {f: int(sum(neto[i][j] for j in idxs)) for f, idxs in por_folio.items()}
        folio_mejor, neto_mejor = max(netos_folio.items(), key=lambda kv: (kv[1], kv[0]))
        if int(comision_total[i]) != 0 and neto_mejor <= 0:
            rut = normalizar_rut_sin_puntos_con_guion(asignaciones[i].rut)
            raise ValueError(
                f"No hay folio elegible para cargar comision del cliente {rut}: "
                f"comision_total={int(comision_total[i])} pero neto=0 en todos los folios."
            )
        folio_por_cliente[i] = int(folio_mejor)
    return folio_por_cliente


def construir_hoja_folio_df(
    folio: int,
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
    x: Sequence[Sequence[int]],
    neto: Sequence[Sequence[int]],
    comision_total: Sequence[int],
    folio_comision_por_cliente: Dict[int, int],
) -> pd.DataFrame:
    idxs = indices_lineas_por_folio(lineas)[folio]
    precios_orden = orden_precios_input_en_folio(lineas, idxs)
    precio_a_idxs = mapa_indices_por_precio_en_folio(lineas, idxs)

    filas: List[Dict[str, object]] = []
    for i, a in enumerate(asignaciones):
        fila: Dict[str, object] = {
            "RUT": normalizar_rut_sin_puntos_con_guion(a.rut),
            "SBR": fondo_como_numero(a.fondo),
            "Tipo": a.tipo,
            "Nemo": a.nemo,
        }
        total_cantidad = 0
        for precio_txt in precios_orden:
            js = precio_a_idxs.get(precio_txt, [])
            q = sum(x[i][j] for j in js)
            fila[precio_txt] = int(q)
            total_cantidad += int(q)

        neto_folio = int(sum(neto[i][j] for j in idxs))
        folio_obj = int(folio_comision_por_cliente.get(i, folio))
        comision_folio = int(comision_total[i]) if folio == folio_obj else 0
        monto_final_folio = int(neto_folio + comision_folio)

        fila["Total"] = int(total_cantidad)
        fila["$ Neto"] = int(neto_folio)
        fila["Comision"] = int(comision_folio)
        fila["Monto Final"] = int(monto_final_folio)
        filas.append(fila)

    df = pd.DataFrame(filas)
    totales: Dict[str, object] = {"RUT": "Totales Acc", "SBR": "", "Tipo": "", "Nemo": ""}
    for precio_txt in precios_orden:
        js = precio_a_idxs.get(precio_txt, [])
        totales[precio_txt] = int(sum(sum(x[i][j] for j in js) for i in range(len(asignaciones))))
    totales["Total"] = int(df["Total"].sum())
    totales["$ Neto"] = int(df["$ Neto"].sum())
    totales["Comision"] = int(df["Comision"].sum())
    totales["Monto Final"] = int(df["Monto Final"].sum())
    return pd.concat([df, pd.DataFrame([totales])], ignore_index=True)


def construir_resumen_df(
    asignaciones: Sequence[Asignacion],
    neto_total: Sequence[int],
    comision_total: Sequence[int],
    folio_comision_por_cliente: Dict[int, int],
) -> pd.DataFrame:
    filas: List[Dict[str, object]] = []
    for i, a in enumerate(asignaciones):
        filas.append(
            {
                "RUT": normalizar_rut_sin_puntos_con_guion(a.rut),
                "SBR": fondo_como_numero(a.fondo),
                "Tipo": a.tipo,
                "Nemo": a.nemo,
                "Cantidad Objetivo": int(a.cantidad_objetivo),
                "Neto Total Output": int(neto_total[i]),
                "Monto Final Asignacion": int(a.monto_final_objetivo),
                "Comision Total": int(comision_total[i]),
                "Folio Comision Cliente": int(folio_comision_por_cliente.get(i, 0)),
            }
        )
    df = pd.DataFrame(filas)
    df_total = pd.DataFrame(
        [
            {
                "RUT": "RESUMEN",
                "SBR": "",
                "Tipo": "",
                "Nemo": "",
                "Cantidad Objetivo": int(df["Cantidad Objetivo"].sum()),
                "Neto Total Output": int(df["Neto Total Output"].sum()),
                "Monto Final Asignacion": int(df["Monto Final Asignacion"].sum()),
                "Comision Total": int(df["Comision Total"].sum()),
                "Folio Comision Cliente": "",
            }
        ]
    )
    return pd.concat([df, df_total], ignore_index=True)


def aplicar_formato_hoja(ws, offset_col: int = 1) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    max_row = ws.max_row
    max_col = ws.max_column
    ws.auto_filter.ref = (
        ws.cell(row=1, column=offset_col).coordinate
        + ":"
        + ws.cell(row=max_row, column=max_col).coordinate
    )

    for r in range(1, max_row + 1):
        for c in range(offset_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if r == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

    headers = [ws.cell(row=1, column=c).value for c in range(offset_col, max_col + 1)]
    for idx, h in enumerate(headers, start=offset_col):
        if h is None:
            continue
        h_str = str(h)
        if h_str in {"Total"} or re.fullmatch(r"\d{1,3},\d{2}", h_str):
            for r in range(2, max_row + 1):
                ws.cell(row=r, column=idx).number_format = "#,##0"
        if h_str in {"$ Neto", "Comision", "Monto Final"}:
            for r in range(2, max_row + 1):
                ws.cell(row=r, column=idx).number_format = "#,##0"

    for c in range(offset_col, max_col + 1):
        letter = get_column_letter(c)
        max_len = 0
        for r in range(1, max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 35)

    if offset_col > 1:
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 12
        ws["A1"].font = Font(bold=True)


def generar_archivos_carga(
    carpeta_salida: Path,
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
    x: Sequence[Sequence[int]],
) -> List[Path]:
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    por_folio = indices_lineas_por_folio(lineas)
    rutas: List[Path] = []

    for folio, idxs in sorted(por_folio.items()):
        precios_orden = orden_precios_input_en_folio(lineas, idxs)
        precio_a_idxs = mapa_indices_por_precio_en_folio(lineas, idxs)

        filas: List[Dict[str, object]] = []
        for i, a in enumerate(asignaciones):
            rut = normalizar_rut_sin_puntos_con_guion(a.rut)
            sbr = fondo_como_numero(a.fondo)
            for precio_txt in precios_orden:
                js = precio_a_idxs.get(precio_txt, [])
                cnt = int(sum(x[i][j] for j in js))
                if cnt > 0:
                    filas.append({"RUT": rut, "SBR": sbr, "Precio": precio_txt, "CNT": cnt})

        df_carga = pd.DataFrame(filas, columns=["RUT", "SBR", "Precio", "CNT"])
        ruta_archivo = carpeta_salida / f"Factura{folio}.xlsx"
        with pd.ExcelWriter(ruta_archivo, engine="openpyxl") as writer:
            df_carga.to_excel(writer, sheet_name="carga", index=False)
        rutas.append(ruta_archivo)

    return rutas


def escribir_output_principal(
    ruta_salida: Path,
    asignaciones: Sequence[Asignacion],
    lineas: Sequence[LineaFactura],
    x: Sequence[Sequence[int]],
    neto: Sequence[Sequence[int]],
    neto_total: Sequence[int],
    comision_total: Sequence[int],
    folio_comision_por_cliente: Dict[int, int],
) -> None:
    folios = sorted({l.folio for l in lineas})
    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        df_resumen = construir_resumen_df(asignaciones, neto_total, comision_total, folio_comision_por_cliente)
        df_resumen.to_excel(writer, sheet_name="resumen", index=False)
        ws_resumen = writer.book["resumen"]
        aplicar_formato_hoja(ws_resumen, offset_col=1)
        ws_resumen.sheet_state = "hidden"

        for folio in folios:
            df_folio = construir_hoja_folio_df(
                folio=folio,
                asignaciones=asignaciones,
                lineas=lineas,
                x=x,
                neto=neto,
                comision_total=comision_total,
                folio_comision_por_cliente=folio_comision_por_cliente,
            )
            nombre_hoja = f"folio_{folio}"[:31]
            df_folio.to_excel(writer, sheet_name=nombre_hoja, index=False, startrow=0, startcol=2)
            ws = writer.book[nombre_hoja]
            ws["A1"] = "FOLIO"
            ws["B1"] = int(folio)
            aplicar_formato_hoja(ws, offset_col=3)


def obtener_nemo_ejecucion(asignaciones: Sequence[Asignacion]) -> str:
    nemo = normalizar_texto(asignaciones[0].nemo) if asignaciones else ""
    for ch in ["/", "\\", ":", "*", "?", "\"", "<", ">", "|"]:
        nemo = nemo.replace(ch, "-")
    nemo = nemo.strip()
    return nemo or "SIN_NEMO"


def ejecutar_proceso(ruta_excel_entrada: Path, carpeta_salida: Path) -> ResultadoProceso:
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    asignaciones, lineas = leer_excel_entrada(ruta_excel_entrada)
    x = construir_asignacion_cantidades(asignaciones, lineas)
    aplicar_swaps_para_signo(asignaciones, lineas, x, max_iteraciones=50_000)
    neto = repartir_neto_por_linea(x, lineas)
    neto_total, comision_total = calcular_neto_y_comision_total(asignaciones, neto)

    malos = violaciones_signo(asignaciones, comision_total)
    if malos:
        ejemplos = []
        for i in malos[:15]:
            rut = normalizar_rut_sin_puntos_con_guion(asignaciones[i].rut)
            fondo = fondo_como_numero(asignaciones[i].fondo)
            tipo = asignaciones[i].tipo
            ejemplos.append(f"{rut} | fondo {fondo} | {tipo} | comision={int(comision_total[i])}")
        raise ValueError("Persisten violaciones de signo:\n" + "\n".join(ejemplos))

    total_neto_facturas = sum(l.monto_neto for l in lineas)
    if sum(neto_total) != total_neto_facturas:
        raise RuntimeError("Error interno: neto total output no cuadra con total neto facturas.")

    folio_comision_por_cliente = elegir_folio_comision_por_cliente(
        asignaciones=asignaciones,
        lineas=lineas,
        neto=neto,
        comision_total=comision_total,
    )

    nemo = obtener_nemo_ejecucion(asignaciones)
    ruta_excel_salida = carpeta_salida / f"Planilla Distribucion {nemo}.xlsx"
    escribir_output_principal(
        ruta_salida=ruta_excel_salida,
        asignaciones=asignaciones,
        lineas=lineas,
        x=x,
        neto=neto,
        neto_total=neto_total,
        comision_total=comision_total,
        folio_comision_por_cliente=folio_comision_por_cliente,
    )

    carpeta_carga = carpeta_salida / "archivos_carga"
    generar_archivos_carga(carpeta_salida=carpeta_carga, asignaciones=asignaciones, lineas=lineas, x=x)

    return ResultadoProceso(
        ruta_planilla=ruta_excel_salida,
        carpeta_archivos_carga=carpeta_carga,
        cantidad_folios=len({l.folio for l in lineas}),
        cantidad_lineas_factura=len(lineas),
        cantidad_asignaciones=len(asignaciones),
    )


def dataframe_facturas_limpio(df_facturas: pd.DataFrame) -> pd.DataFrame:
    _validar_columnas(df_facturas, COLUMNAS_FACTURAS, "facturas")
    return df_facturas.loc[:, COLUMNAS_FACTURAS].copy()


def crear_excel_con_facturas_reemplazadas(
    ruta_excel_base: Path,
    carpeta_salida: Path,
    df_facturas: pd.DataFrame,
    nombre_archivo: str | None = None,
) -> Path:
    ruta_excel_base = Path(ruta_excel_base)
    carpeta_salida = Path(carpeta_salida)
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    if not ruta_excel_base.exists():
        raise FileNotFoundError(f"No existe el Excel base: {ruta_excel_base}")

    df_facturas_out = dataframe_facturas_limpio(df_facturas)
    xls = pd.ExcelFile(ruta_excel_base)
    if "asignaciones" not in xls.sheet_names:
        raise ValueError("El Excel base no contiene hoja 'asignaciones'.")

    hojas: Dict[str, pd.DataFrame] = {}
    for hoja in xls.sheet_names:
        if hoja == "facturas":
            continue
        hojas[hoja] = pd.read_excel(ruta_excel_base, sheet_name=hoja, dtype=object)

    hojas["facturas"] = df_facturas_out
    if nombre_archivo is None:
        nombre_archivo = f"{ruta_excel_base.stem}_facturas_pdf.xlsx"
    ruta_salida = carpeta_salida / nombre_archivo

    orden_hojas = []
    for hoja in xls.sheet_names:
        if hoja == "facturas":
            orden_hojas.append("facturas")
        else:
            orden_hojas.append(hoja)
    if "facturas" not in orden_hojas:
        orden_hojas.append("facturas")

    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        escritos = set()
        for hoja in orden_hojas:
            if hoja in escritos or hoja not in hojas:
                continue
            hojas[hoja].to_excel(writer, sheet_name=hoja, index=False)
            escritos.add(hoja)
        for hoja, df in hojas.items():
            if hoja in escritos:
                continue
            df.to_excel(writer, sheet_name=hoja[:31], index=False)

    return ruta_salida

