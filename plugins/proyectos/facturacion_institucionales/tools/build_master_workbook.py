# -*- coding: utf-8 -*-
"""
Created on Thu Jan 23 10:24:12 2025

@author: BBRUNA
"""

from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    Alignment = Font = PatternFill = None
    get_column_letter = None


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PLUGIN_ROOT / "Facturacion_Institucionales_Maestro.xlsx"
SOURCE_BASENAME = "CONFIG Facturacion Institucionales.xlsx"


CLIENT_SOURCE_COLUMNS = {
    "Prioridad": "prioridad",
    "RUT": "rut_origen",
    "CLIENTE": "cliente_nombre",
    "Operador\nDirecto": "operador_directo_origen",
    "FACTURA": "factura_origen",
    " Resumen Transacciones\nPDF": "resumen_origen",
    "Reporte Excel\nProcesos de Datos": "reporte_proceso_origen",
    "Archivo Plano (si es por OMS)": "archivo_plano_origen",
    "Grupo": "grupo_nombre",
    "Carpeta": "carpeta_salida",
    "Clave (Si es vacia se dejan los ultimos numeros del rut)": "clave_override_origen",
    "TamaÃ±o maximo (MB)": "tamano_max_mb_override_origen",
    "Columna Emisor para Reporte Excel Proceso de Datos": "reporte_emisor_columna",
    "Fila Inicio para Reporte Excel Proceso de Datos": "reporte_fila_inicio_origen",
    "Observaciones": "observaciones",
}

MAIL_SOURCE_COLUMNS = {
    "Grupo": "grupo_nombre",
    "Email Destinatarios": "emails_origen",
    "Clave Archivos": "clave_grupo_origen",
    "TamaÃ±o MÃ¡ximo (MB)": "tamano_max_mb_grupo_origen",
}

MAIL_REGEX = re.compile(r"[A-Za-z0-9._%+\-']+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
SQL_START_REGEX = re.compile(r"^\s*(select|with)\b", re.IGNORECASE | re.DOTALL)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _norm_key(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value).casefold()).strip()


def _norm_rut(value: Any) -> str:
    return re.sub(r"[^0-9K]", "", _text(value).upper())


def _rut_parts(rut_norm: str) -> tuple[str, str]:
    if not rut_norm:
        return "", ""
    if len(rut_norm) == 1:
        return "", rut_norm
    return rut_norm[:-1], rut_norm[-1]


def _to_int(value: Any) -> int | None:
    raw = _text(value)
    if not raw:
        return None
    try:
        return int(float(raw))
    except Exception:
        return None


def _to_float(value: Any) -> float | None:
    raw = _text(value)
    if not raw:
        return None
    try:
        return float(raw)
    except Exception:
        return None


def _flag_yn(value: Any, field: str) -> tuple[str, str]:
    token = _text(value).upper()
    if token in {"", "N"}:
        return "N", ""
    if token == "Y":
        return "Y", ""
    return "N", f"Valor no estandar en {field}: '{token}'"


def _parse_factura_field(value: Any) -> tuple[str, str, str]:
    token = _text(value).upper()
    if token in {"", "N"}:
        return "N", "N", ""
    if token in {"Y", "PDF"}:
        return "Y", "N", ""
    if token == "XML":
        return "Y", "Y", ""
    return "N", "N", f"Valor FACTURA no estandar: '{token}'"


@dataclass
class SqlFieldDecision:
    estado: str
    sql: str
    aviso: str


def _parse_sql_field(value: Any) -> SqlFieldDecision:
    raw = _text(value)
    token = raw.upper()
    if not raw or token == "N":
        return SqlFieldDecision("NO", "", "")
    if token == "Y":
        return SqlFieldDecision(
            "PENDIENTE_QUERY",
            "",
            "Requiere query, pero aun no esta definida en el archivo de origen.",
        )
    if SQL_START_REGEX.match(raw):
        return SqlFieldDecision("QUERY", raw, "")
    return SqlFieldDecision(
        "PENDIENTE_REVISAR",
        raw,
        "Valor no estandar: revisar si corresponde query valida.",
    )


def _split_emails(raw: str) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for match in MAIL_REGEX.finditer(raw or ""):
        email = match.group(0).strip(" ,;")
        key = email.casefold()
        if key in seen:
            continue
        seen.add(key)
        result.append(email)
    return result


def _rename_keep(df: pd.DataFrame, mapping: dict[str, str]) -> pd.DataFrame:
    renamed = df.rename(columns=mapping).copy()
    for col in mapping.values():
        if col not in renamed.columns:
            renamed[col] = ""
    return renamed[list(mapping.values())]


def _issue(severity: str, sheet: str, record: str, field: str, detail: str) -> dict[str, Any]:
    return {
        "severity": severity,
        "sheet": sheet,
        "record_key": record,
        "field": field,
        "detail": detail,
    }


def _find_source_file(explicit: str | None) -> Path:
    if explicit:
        src = Path(explicit).expanduser()
        if not src.is_absolute():
            src = (Path.cwd() / src).resolve()
        if not src.exists():
            raise FileNotFoundError(f"No existe archivo fuente: {src}")
        return src

    candidates = sorted(Path.cwd().rglob(SOURCE_BASENAME))
    if not candidates:
        raise FileNotFoundError(
            f"No se encontro '{SOURCE_BASENAME}'. Usa --source para indicar la ruta."
        )
    if len(candidates) > 1:
        # Prioriza el archivo mas reciente para evitar tomar copias antiguas por error.
        candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def _build_mail_tables(
    mail_source: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, dict[str, Any]], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    group_rows: list[dict[str, Any]] = []
    email_rows: list[dict[str, Any]] = []
    lookup: dict[str, dict[str, Any]] = {}

    df = _rename_keep(mail_source, MAIL_SOURCE_COLUMNS)
    for idx, row in df.iterrows():
        source_row = idx + 2
        group_name = _text(row["grupo_nombre"])
        group_key = _norm_key(group_name)
        emails_raw = _text(row["emails_origen"])
        default_key = _text(row["clave_grupo_origen"])
        max_mb = _to_float(row["tamano_max_mb_grupo_origen"])
        emails = _split_emails(emails_raw)

        row_record = {
            "source_row_num": source_row,
            "grupo_nombre": group_name,
            "grupo_key": group_key,
            "correo_manual_habilitado": "Y" if emails else "N",
            "emails_origen": emails_raw,
            "cantidad_emails_detectados": len(emails),
            "clave_archivos_grupo": default_key,
            "tamano_max_mb_grupo": max_mb,
            "nota_migracion": "",
        }

        if not group_key:
            issues.append(_issue("ERROR", "CORREOS_GRUPO", f"row:{source_row}", "grupo_nombre", "Grupo vacio."))
        elif group_key in lookup:
            issues.append(
                _issue(
                    "WARN",
                    "CORREOS_GRUPO",
                    group_name or f"row:{source_row}",
                    "grupo_nombre",
                    "Grupo duplicado. Se conserva la primera fila para lookup.",
                )
            )
            row_record["nota_migracion"] = "Duplicado en origen."
        else:
            lookup[group_key] = row_record

        if emails_raw and not emails:
            issues.append(
                _issue(
                    "WARN",
                    "CORREOS_GRUPO",
                    group_name or f"row:{source_row}",
                    "emails_origen",
                    "No se detectaron correos validos en el texto.",
                )
            )

        for order, email in enumerate(emails, start=1):
            email_rows.append(
                {
                    "grupo_nombre": group_name,
                    "grupo_key": group_key,
                    "orden_email": order,
                    "email": email,
                    "email_normalizado": email.casefold(),
                    "source_row_num": source_row,
                }
            )

        group_rows.append(row_record)

    group_df = pd.DataFrame(group_rows).sort_values(by=["grupo_nombre", "source_row_num"], na_position="last")
    email_df = pd.DataFrame(email_rows).sort_values(by=["grupo_nombre", "orden_email"], na_position="last")
    return group_df, email_df, lookup, issues


def _transform_membership_sheet(
    raw_sheet: pd.DataFrame,
    source_label: str,
    mail_lookup: dict[str, dict[str, Any]],
) -> tuple[pd.DataFrame, pd.DataFrame, list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    query_rows: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []

    df = _rename_keep(raw_sheet, CLIENT_SOURCE_COLUMNS)
    prefix = "CLI" if source_label == "CLIENTES" else "GRP"
    query_counter = 0

    for idx, row in df.iterrows():
        source_row = idx + 2
        record_id = f"{prefix}{idx + 1:04d}"

        rut_raw = _text(row["rut_origen"])
        rut_norm = _norm_rut(rut_raw)
        rut_num, rut_dv = _rut_parts(rut_norm)
        rut_last4 = rut_num[-4:] if rut_num else ""

        client_name = _text(row["cliente_nombre"])
        group_name = _text(row["grupo_nombre"])
        group_key = _norm_key(group_name)
        folder_name = _text(row["carpeta_salida"])

        priority = _to_int(row["prioridad"])
        direct_operator, warn = _flag_yn(row["operador_directo_origen"], "operador_directo")
        if warn:
            issues.append(_issue("WARN", source_label, record_id, "operador_directo", warn))

        pdf_flag, xml_flag, warn = _parse_factura_field(row["factura_origen"])
        if warn:
            issues.append(_issue("WARN", source_label, record_id, "factura_origen", warn))

        summary_flag, warn = _flag_yn(row["resumen_origen"], "resumen_pdf")
        if warn:
            issues.append(_issue("WARN", source_label, record_id, "resumen_pdf", warn))

        report_decision = _parse_sql_field(row["reporte_proceso_origen"])
        flat_decision = _parse_sql_field(row["archivo_plano_origen"])
        if report_decision.aviso:
            issues.append(_issue("INFO", source_label, record_id, "reporte_proceso", report_decision.aviso))
        if flat_decision.aviso:
            issues.append(_issue("INFO", source_label, record_id, "archivo_plano", flat_decision.aviso))

        report_query_id = ""
        flat_query_id = ""
        report_emisor_col = _text(row["reporte_emisor_columna"])
        report_start_row = _to_int(row["reporte_fila_inicio_origen"])

        if report_decision.estado in {"QUERY", "PENDIENTE_REVISAR"} and report_decision.sql:
            query_counter += 1
            report_query_id = f"{prefix}_REP_{query_counter:04d}"
            query_rows.append(
                {
                    "query_id": report_query_id,
                    "source_sheet": source_label,
                    "source_record_id": record_id,
                    "source_row_num": source_row,
                    "query_kind": "REPORTE_PROCESO_DATOS",
                    "query_status": report_decision.estado,
                    "rut_norm": rut_norm,
                    "cliente_nombre": client_name,
                    "grupo_nombre": group_name,
                    "sql_text": report_decision.sql,
                    "emisor_columna": report_emisor_col,
                    "fila_inicio": report_start_row,
                }
            )

        if flat_decision.estado in {"QUERY", "PENDIENTE_REVISAR"} and flat_decision.sql:
            query_counter += 1
            flat_query_id = f"{prefix}_PLN_{query_counter:04d}"
            query_rows.append(
                {
                    "query_id": flat_query_id,
                    "source_sheet": source_label,
                    "source_record_id": record_id,
                    "source_row_num": source_row,
                    "query_kind": "ARCHIVO_PLANO_OMS",
                    "query_status": flat_decision.estado,
                    "rut_norm": rut_norm,
                    "cliente_nombre": client_name,
                    "grupo_nombre": group_name,
                    "sql_text": flat_decision.sql,
                    "emisor_columna": "",
                    "fila_inicio": None,
                }
            )

        mail_match = mail_lookup.get(group_key)
        mail_status = "NO_MATCH"
        mail_lookup_by = ""
        mail_group_name = ""
        mail_enabled = "N"
        group_default_key = ""
        group_default_max_mb = None
        if mail_match:
            mail_status = "MATCH"
            mail_lookup_by = "GRUPO"
            mail_group_name = str(mail_match.get("grupo_nombre", "") or group_name)
            mail_enabled = str(mail_match.get("correo_manual_habilitado", "N"))
            group_default_key = _text(mail_match.get("clave_archivos_grupo", ""))
            group_default_max_mb = mail_match.get("tamano_max_mb_grupo")
        elif group_name:
            issues.append(
                _issue(
                    "WARN",
                    source_label,
                    record_id,
                    "grupo_nombre",
                    f"No existe grupo en CORREOS para '{group_name}'.",
                )
            )

        client_key_override = _text(row["clave_override_origen"])
        if client_key_override:
            final_key = client_key_override
            key_source = "CLIENTE_OVERRIDE"
        elif group_default_key:
            final_key = group_default_key
            key_source = "CORREO_GRUPO"
        else:
            final_key = rut_last4
            key_source = "RUT_LAST4"

        client_max_mb_override = _to_float(row["tamano_max_mb_override_origen"])
        final_max_mb = (
            client_max_mb_override if client_max_mb_override is not None else group_default_max_mb
        )

        if not rut_norm:
            issues.append(_issue("ERROR", source_label, record_id, "rut_origen", "RUT vacio o invalido."))
        if not client_name:
            issues.append(_issue("WARN", source_label, record_id, "cliente_nombre", "Cliente vacio."))

        records.append(
            {
                "record_id": record_id,
                "source_row_num": source_row,
                "activo": "Y",
                "prioridad": priority,
                "rut_origen": rut_raw,
                "rut_norm": rut_norm,
                "rut_sin_dv": rut_num,
                "rut_dv": rut_dv,
                "cliente_nombre": client_name,
                "grupo_nombre": group_name,
                "carpeta_salida": folder_name,
                "operador_directo": direct_operator,
                "factura_pdf": pdf_flag,
                "factura_xml": xml_flag,
                "resumen_pdf": summary_flag,
                "reporte_proceso_estado": report_decision.estado,
                "reporte_proceso_query_id": report_query_id,
                "reporte_proceso_emisor_columna": report_emisor_col,
                "reporte_proceso_fila_inicio": report_start_row,
                "archivo_plano_estado": flat_decision.estado,
                "archivo_plano_query_id": flat_query_id,
                "correo_manual_habilitado": mail_enabled,
                "correo_lookup_status": mail_status,
                "correo_lookup_by": mail_lookup_by,
                "correo_grupo": mail_group_name,
                "clave_archivos_override": client_key_override,
                "clave_archivos_grupo": group_default_key,
                "clave_archivos_efectiva": final_key,
                "clave_archivos_fuente": key_source,
                "clave_archivos_rut_last4": rut_last4,
                "tamano_max_mb_override": client_max_mb_override,
                "tamano_max_mb_grupo": group_default_max_mb,
                "tamano_max_mb_efectivo": final_max_mb,
                "observaciones": _text(row["observaciones"]),
                "factura_origen": _text(row["factura_origen"]),
                "resumen_origen": _text(row["resumen_origen"]),
                "reporte_proceso_origen": _text(row["reporte_proceso_origen"]),
                "archivo_plano_origen": _text(row["archivo_plano_origen"]),
            }
        )

    normalized_df = pd.DataFrame(records)
    queries_df = pd.DataFrame(query_rows)
    return normalized_df, queries_df, issues


def _add_duplicate_rut_issues(df: pd.DataFrame, sheet: str, issues: list[dict[str, Any]]) -> None:
    if df.empty:
        return
    duplicates = df[df["rut_norm"].astype(str).duplicated(keep=False) & df["rut_norm"].astype(str).ne("")]
    if duplicates.empty:
        return
    for rut, group in duplicates.groupby("rut_norm"):
        record_list = ", ".join(group["record_id"].astype(str).tolist())
        issues.append(
            _issue("WARN", sheet, rut, "rut_norm", f"RUT duplicado en hoja normalizada. Registros: {record_list}")
        )


def _build_rules_sheet() -> pd.DataFrame:
    rows = [
        ("FACTURA", "Y", "factura_pdf=Y, factura_xml=N", "Descarga solo factura PDF"),
        ("FACTURA", "XML", "factura_pdf=Y, factura_xml=Y", "Descarga factura PDF y XML"),
        ("Resumen Transacciones PDF", "Y", "resumen_pdf=Y", "Genera resumen en PDF"),
        (
            "Reporte Excel Procesos de Datos",
            "Y",
            "reporte_proceso_estado=PENDIENTE_QUERY",
            "Mostrar aviso de query pendiente al terminar el resto del flujo",
        ),
        (
            "Reporte Excel Procesos de Datos",
            "SELECT.../WITH...",
            "reporte_proceso_estado=QUERY + fila en QUERIES",
            "Ejecuta consulta y exporta Excel",
        ),
        (
            "Archivo Plano (si es por OMS)",
            "Y",
            "archivo_plano_estado=PENDIENTE_QUERY",
            "Mostrar aviso de query pendiente",
        ),
        (
            "Archivo Plano (si es por OMS)",
            "SELECT.../WITH...",
            "archivo_plano_estado=QUERY + fila en QUERIES",
            "Ejecuta consulta y exporta archivo plano",
        ),
        (
            "Clave de archivos",
            "override / grupo / rut",
            "CLIENTE_OVERRIDE > CORREO_GRUPO > RUT_LAST4",
            "Resolucion de clave efectiva para PDFs cifrados",
        ),
    ]
    return pd.DataFrame(rows, columns=["campo_origen", "valor_origen", "resultado_maestro", "significado"])


def _build_readme_sheet(source_name: str, output_name: str) -> pd.DataFrame:
    rows = [
        ("GENERAL", "archivo", f"{output_name} - libro maestro de Facturacion Institucionales."),
        ("GENERAL", "fuente", f"Generado desde {source_name} (hojas CLIENTES, GRUPOS, CORREOS)."),
        ("HOJA", "CLIENTES", "Registros individuales normalizados y listos para ejecucion."),
        ("HOJA", "GRUPOS_MIEMBROS", "Miembros provenientes de hoja GRUPOS, mismo esquema operativo."),
        ("HOJA", "QUERIES", "Consultas SQL extraidas de columnas mixtas."),
        ("HOJA", "CORREOS_GRUPO", "Configuracion de correos y valores por grupo."),
        ("HOJA", "CORREOS_EMAILS", "Un email por fila para validacion y mantenimiento."),
        ("HOJA", "REGLAS", "Reglas de transformacion y semantica del maestro."),
        ("HOJA", "VALIDACION", "Hallazgos de calidad durante la migracion."),
        ("REGLA", "operacion_factura", "Y=PDF, XML=PDF+XML."),
        ("REGLA", "estado_query", "NO / PENDIENTE_QUERY / QUERY / PENDIENTE_REVISAR."),
    ]
    return pd.DataFrame(rows, columns=["seccion", "item", "detalle"])


def _build_meta_sheet(
    *,
    source_file: Path,
    output_file: Path,
    clientes_df: pd.DataFrame,
    grupos_df: pd.DataFrame,
    queries_df: pd.DataFrame,
    correos_df: pd.DataFrame,
    emails_df: pd.DataFrame,
    issues_df: pd.DataFrame,
) -> pd.DataFrame:
    issue_counts = issues_df["severity"].value_counts().to_dict() if not issues_df.empty else {}
    rows = [
        ("source_file", source_file.name),
        ("output_file", output_file.name),
        ("clientes_rows", len(clientes_df)),
        ("grupos_miembros_rows", len(grupos_df)),
        ("queries_rows", len(queries_df)),
        ("correos_grupo_rows", len(correos_df)),
        ("correos_emails_rows", len(emails_df)),
        ("validation_rows", len(issues_df)),
        ("validation_errors", int(issue_counts.get("ERROR", 0))),
        ("validation_warn", int(issue_counts.get("WARN", 0))),
        ("validation_info", int(issue_counts.get("INFO", 0))),
    ]
    return pd.DataFrame(rows, columns=["key", "value"])


def _style_workbook(path: Path) -> None:
    if get_column_letter is None:
        return
    try:
        from openpyxl import load_workbook
    except Exception:
        return

    wb = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78") if PatternFill else None
    header_font = Font(color="FFFFFF", bold=True) if Font else None
    wrapped = Alignment(wrap_text=True, vertical="top") if Alignment else None

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            if header_fill:
                cell.fill = header_fill
            if header_font:
                cell.font = header_font
            if wrapped:
                cell.alignment = wrapped

        for col_idx, cells in enumerate(ws.columns, start=1):
            max_len = 0
            sample = list(cells)[:1000]
            for cell in sample:
                value = "" if cell.value is None else str(cell.value)
                longest_line = max((len(part) for part in value.splitlines()), default=0)
                if longest_line > max_len:
                    max_len = longest_line
            width = min(max(12, max_len + 2), 90)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        wrap_headers = {
            "sql_text",
            "emails_origen",
            "detalle",
            "observaciones",
            "reporte_proceso_origen",
            "archivo_plano_origen",
        }
        header_map = {str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}
        for header in wrap_headers:
            column_index = header_map.get(header)
            if not column_index:
                continue
            for row_idx in range(2, min(ws.max_row, 5000) + 1):
                cell = ws.cell(row=row_idx, column=column_index)
                if wrapped:
                    cell.alignment = wrapped

    wb.save(path)


def build_master_workbook(source_file: Path, output_file: Path) -> dict[str, Any]:
    excel = pd.ExcelFile(source_file)
    expected = {"CLIENTES", "GRUPOS", "CORREOS"}
    missing = expected - set(excel.sheet_names)
    if missing:
        raise ValueError(f"Faltan hojas requeridas: {sorted(missing)}")

    source_clients = pd.read_excel(source_file, sheet_name="CLIENTES")
    source_groups = pd.read_excel(source_file, sheet_name="GRUPOS")
    source_mails = pd.read_excel(source_file, sheet_name="CORREOS")

    mail_group_df, mail_email_df, mail_lookup, issues = _build_mail_tables(source_mails)
    clients_df, client_queries_df, client_issues = _transform_membership_sheet(
        source_clients, "CLIENTES", mail_lookup
    )
    group_members_df, group_queries_df, group_issues = _transform_membership_sheet(
        source_groups, "GRUPOS_MIEMBROS", mail_lookup
    )
    issues.extend(client_issues)
    issues.extend(group_issues)

    queries_df = pd.concat([client_queries_df, group_queries_df], ignore_index=True)
    _add_duplicate_rut_issues(clients_df, "CLIENTES", issues)
    _add_duplicate_rut_issues(group_members_df, "GRUPOS_MIEMBROS", issues)

    if not clients_df.empty and not group_members_df.empty:
        in_clients = set(clients_df["rut_norm"].astype(str)) - {""}
        in_groups = set(group_members_df["rut_norm"].astype(str)) - {""}
        for rut in sorted(in_clients & in_groups):
            issues.append(
                _issue(
                    "INFO",
                    "CROSSCHECK",
                    rut,
                    "rut_norm",
                    "RUT presente en CLIENTES y GRUPOS_MIEMBROS; revisar si corresponde doble flujo.",
                )
            )

    issues_df = pd.DataFrame(issues)
    if issues_df.empty:
        issues_df = pd.DataFrame(columns=["severity", "sheet", "record_key", "field", "detail"])
    else:
        issues_df = issues_df.sort_values(by=["severity", "sheet", "record_key", "field"])

    operational_columns = [
        "record_id",
        "source_row_num",
        "activo",
        "prioridad",
        "rut_origen",
        "rut_norm",
        "rut_sin_dv",
        "rut_dv",
        "cliente_nombre",
        "grupo_nombre",
        "carpeta_salida",
        "operador_directo",
        "factura_pdf",
        "factura_xml",
        "resumen_pdf",
        "reporte_proceso_estado",
        "reporte_proceso_query_id",
        "reporte_proceso_emisor_columna",
        "reporte_proceso_fila_inicio",
        "archivo_plano_estado",
        "archivo_plano_query_id",
        "correo_manual_habilitado",
        "correo_lookup_status",
        "correo_lookup_by",
        "correo_grupo",
        "clave_archivos_override",
        "clave_archivos_grupo",
        "clave_archivos_efectiva",
        "clave_archivos_fuente",
        "clave_archivos_rut_last4",
        "tamano_max_mb_override",
        "tamano_max_mb_grupo",
        "tamano_max_mb_efectivo",
        "observaciones",
        "factura_origen",
        "resumen_origen",
        "reporte_proceso_origen",
        "archivo_plano_origen",
    ]
    clients_df = clients_df[operational_columns]
    group_members_df = group_members_df[operational_columns]

    query_columns = [
        "query_id",
        "source_sheet",
        "source_record_id",
        "source_row_num",
        "query_kind",
        "query_status",
        "rut_norm",
        "cliente_nombre",
        "grupo_nombre",
        "sql_text",
        "emisor_columna",
        "fila_inicio",
    ]
    if queries_df.empty:
        queries_df = pd.DataFrame(columns=query_columns)
    else:
        queries_df = queries_df[query_columns].sort_values(by=["source_sheet", "source_row_num", "query_kind"])

    readme_df = _build_readme_sheet(source_file.name, output_file.name)
    rules_df = _build_rules_sheet()
    meta_df = _build_meta_sheet(
        source_file=source_file,
        output_file=output_file,
        clientes_df=clients_df,
        grupos_df=group_members_df,
        queries_df=queries_df,
        correos_df=mail_group_df,
        emails_df=mail_email_df,
        issues_df=issues_df,
    )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        readme_df.to_excel(writer, sheet_name="README", index=False)
        meta_df.to_excel(writer, sheet_name="META", index=False)
        clients_df.to_excel(writer, sheet_name="CLIENTES", index=False)
        group_members_df.to_excel(writer, sheet_name="GRUPOS_MIEMBROS", index=False)
        queries_df.to_excel(writer, sheet_name="QUERIES", index=False)
        mail_group_df.to_excel(writer, sheet_name="CORREOS_GRUPO", index=False)
        mail_email_df.to_excel(writer, sheet_name="CORREOS_EMAILS", index=False)
        rules_df.to_excel(writer, sheet_name="REGLAS", index=False)
        issues_df.to_excel(writer, sheet_name="VALIDACION", index=False)

    _style_workbook(output_file)

    return {
        "output": str(output_file),
        "clientes_rows": int(len(clients_df)),
        "grupos_miembros_rows": int(len(group_members_df)),
        "queries_rows": int(len(queries_df)),
        "correos_grupo_rows": int(len(mail_group_df)),
        "correos_emails_rows": int(len(mail_email_df)),
        "validation_rows": int(len(issues_df)),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Construye el workbook maestro de Facturacion Institucionales.")
    parser.add_argument("--source", help="Ruta al archivo CONFIG Facturacion Institucionales.xlsx")
    parser.add_argument("--output", help="Ruta de salida del workbook maestro (.xlsx)")
    args = parser.parse_args()

    source_file = _find_source_file(args.source)
    output_file = Path(args.output).expanduser().resolve() if args.output else DEFAULT_OUTPUT
    result = build_master_workbook(source_file.resolve(), output_file)

    print("Workbook maestro generado:")
    for key, value in result.items():
        print(f" - {key}: {value}")


if __name__ == "__main__":
    main()
